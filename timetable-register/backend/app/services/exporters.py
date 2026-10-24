"""PDF, Excel and CSV exports of timetable views."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..scheduler.timegrid import DAY_NAMES


def _tint(hex_color: str, factor: float = 0.82) -> str:
    h = hex_color.lstrip("#")
    try:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return "#eef2ff"
    r, g, b = (int(c + (255 - c) * factor) for c in (r, g, b))
    return f"#{r:02x}{g:02x}{b:02x}"


def _cell_entries(view: dict, day: int, slot: dict, row_id: int | None = None) -> list[dict]:
    out = []
    for e in view["entries"]:
        if e["day"] != day or not (e["start"] < slot["end"] and slot["start"] < e["end"]):
            continue
        if row_id is not None and e["class_id"] != row_id:
            continue
        out.append(e)
    return out


def _cell_text(view: dict, e: dict, sep: str = "\n") -> str:
    parts = [f"{e['subject_code'] or e['subject_name']}"]
    if view["view"] != "class" and view["kind"] == "grid":
        parts.append(e["class_name"])
    if view["view"] != "teacher":
        parts.append(e["teacher_name"])
    if view["view"] not in ("room", "lab"):
        parts.append(e["room_name"])
    return sep.join(p for p in parts if p)


def grid_matrix(view: dict) -> tuple[list[str], list[list[tuple[str, str | None]]]]:
    """Header + rows of (text, color) for a grid (class/teacher/room) view."""
    header = ["Time"] + [d["name"] for d in view["days"]]
    rows = []
    for slot in view["slots"]:
        row: list[tuple[str, str | None]] = [(f"{slot['label']}\n{slot['time']}", None)]
        for d in view["days"]:
            if slot["kind"] != "period":
                present = slot["key"] in view["day_slots"].get(d["day"], [])
                row.append((slot["label"] if present else "", "#f3f4f6"))
                continue
            cell = _cell_entries(view, d["day"], slot)
            if not cell:
                present = slot["key"] in view["day_slots"].get(d["day"], [])
                row.append(("" if present else "—", None))
            else:
                row.append(("\n\n".join(_cell_text(view, e) for e in cell), cell[0]["color"]))
        rows.append(row)
    return header, rows


def master_matrices(view: dict) -> list[tuple[str, list[str], list[list[tuple[str, str | None]]]]]:
    """Per day: rows = classes, columns = periods. Classes on different period grids
    (e.g. 60- vs 50-minute lectures) get separate tables so columns stay meaningful."""
    from ..scheduler.timegrid import fmt_range

    out = []
    for d in view["days"]:
        key = str(d["day"])
        groups: dict[tuple, list[dict]] = {}
        for r in view["rows"]:
            sig = tuple(tuple(p) for p in r.get("periods", {}).get(key, []))
            if sig:
                groups.setdefault(sig, []).append(r)
        multi = len(groups) > 1
        for gi, (sig, rows_in) in enumerate(groups.items(), start=1):
            header = ["Class"] + [f"P{i + 1}\n{fmt_range(s, e)}" for i, (s, e) in enumerate(sig)]
            rows = []
            for r in rows_in:
                row: list[tuple[str, str | None]] = [(r["name"], None)]
                for s, e in sig:
                    slot = {"start": s, "end": e}
                    cell = _cell_entries(view, d["day"], slot, r["id"])
                    if cell:
                        x = cell[0]
                        txt = f"{x['subject_code'] or x['subject_name']}\n{x['teacher_name']}\n{x['room_name']}"
                        if x["start"] < s:
                            txt = f"{x['subject_code']} (contd.)"
                        row.append((txt, x["color"]))
                    else:
                        row.append(("", None))
                rows.append(row)
            label = d["name"] + (f" (period grid {gi})" if multi else "")
            out.append((label, header, rows))
    return out


# ------------------------------------------------------------------- PDF
def to_pdf(views: list[dict], institution: str, timetable_name: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
        title=f"{timetable_name} – {institution}",
    )
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Heading2"], spaceAfter=2)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=6.6, leading=8)
    story = []
    avail_w = landscape(A4)[0] - 20 * mm

    white = ParagraphStyle("w", parent=cell, textColor=colors.white)

    def table_for(header, rows, first_w):
        n = len(header) - 1
        col_w = [first_w] + [(avail_w - first_w) / max(1, n)] * n
        data = [[Paragraph(f"<b>{x.replace(chr(10), '<br/>')}</b>", white) for x in header]]
        style = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c7c9d1")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#f9fafb")),
        ]
        for ri, row in enumerate(rows, start=1):
            prow = []
            for ci, (text, color) in enumerate(row):
                prow.append(Paragraph(text.replace("&", "&amp;").replace("\n", "<br/>"), cell))
                if color and ci > 0:
                    fill = color if color == "#f3f4f6" else _tint(color)
                    style.append(("BACKGROUND", (ci, ri), (ci, ri), colors.HexColor(fill)))
            data.append(prow)
        t = Table(data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle(style))
        return t

    for i, view in enumerate(views):
        if view["kind"] == "grid":
            story.append(Paragraph(f"{view['title']}", h))
            story.append(Paragraph(f"{view.get('subtitle', '')} · {timetable_name} · {institution}", sub))
            story.append(Spacer(1, 4 * mm))
            header, rows = grid_matrix(view)
            story.append(table_for(header, rows, 24 * mm))
            if i < len(views) - 1:
                story.append(PageBreak())
        else:
            mats = master_matrices(view)
            for j, (day, header, rows) in enumerate(mats):
                story.append(Paragraph(f"{view['title']} — {day}", h))
                story.append(Paragraph(f"{timetable_name} · {institution}", sub))
                story.append(Spacer(1, 4 * mm))
                story.append(table_for(header, rows, 32 * mm))
                if j < len(mats) - 1 or i < len(views) - 1:
                    story.append(PageBreak())
    if not story:
        story.append(Paragraph("No sessions to show.", styles["Normal"]))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawString(10 * mm, 6 * mm, f"Generated {datetime.now():%d %b %Y %H:%M} · Timetable Register")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ------------------------------------------------------------------ Excel
def _safe_sheet(name: str, used: set[str]) -> str:
    base = "".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "Sheet"
    out, n = base, 2
    while out in used:
        out = f"{base[:25]}~{n}"
        n += 1
    used.add(out)
    return out


def to_xlsx(views: list[dict], institution: str, timetable_name: str, flat_entries: list[dict]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    thin = Side(style="thin", color="C7C9D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F2937")
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")

    def write(ws, title, subtitle, header, rows, first_w):
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = subtitle
        ws["A2"].font = Font(color="666666", size=9)
        for ci, hname in enumerate(header, start=1):
            c = ws.cell(row=4, column=ci, value=hname)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = wrap
            c.border = border
        for ri, row in enumerate(rows, start=5):
            ws.row_dimensions[ri].height = 48
            for ci, (text, color) in enumerate(row, start=1):
                c = ws.cell(row=ri, column=ci, value=text)
                c.alignment = wrap
                c.border = border
                if color and ci > 1:
                    c.fill = PatternFill("solid", fgColor=(_tint(color) if color != "#f3f4f6" else color).lstrip("#").upper())
                if ci == 1:
                    c.font = Font(bold=True)
        ws.column_dimensions["A"].width = first_w
        for ci in range(2, len(header) + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 22
        ws.freeze_panes = "B5"

    for view in views:
        if view["kind"] == "grid":
            header, rows = grid_matrix(view)
            ws = wb.create_sheet(_safe_sheet(view["title"], used))
            write(ws, view["title"], f"{view.get('subtitle', '')} · {timetable_name} · {institution}", header, rows, 16)
        else:
            for day, header, rows in master_matrices(view):
                ws = wb.create_sheet(_safe_sheet(f"{day[:3]} {view['title']}", used))
                write(ws, f"{view['title']} — {day}", f"{timetable_name} · {institution}", header, rows, 26)

    ws = wb.create_sheet(_safe_sheet("All sessions", used))
    cols = ["Day", "Start", "End", "Class", "Subject code", "Subject", "Teacher", "Room", "Periods", "Locked"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for e in flat_entries:
        s, en = e["time"].split("–")
        ws.append([e["day_name"], s, en, e["class_name"], e["subject_code"], e["subject_name"],
                   e["teacher_name"], e["room_name"], e["length"], "yes" if e["locked"] else ""])
    for i, w in enumerate([12, 8, 8, 26, 12, 28, 24, 18, 8, 8], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -------------------------------------------------------------------- CSV
def to_csv(flat_entries: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["day", "day_name", "start", "end", "period", "periods", "class", "subject_code", "subject",
                "teacher", "room", "room_type", "locked"])
    for e in sorted(flat_entries, key=lambda x: (x["day"], x["start"], x["class_name"])):
        s, en = e["time"].split("–")
        w.writerow([e["day"], DAY_NAMES[e["day"]], s, en, e["period_index"] + 1, e["length"], e["class_name"],
                    e["subject_code"], e["subject_name"], e["teacher_name"], e["room_name"], e["room_type"],
                    int(e["locked"])])
    return ("﻿" + buf.getvalue()).encode("utf-8")
